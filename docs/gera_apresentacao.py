"""Gera a apresentacao do projeto em PowerPoint, com a identidade do Vitrine Design System.

POR QUE o deck e um script e nao um .pptx editado a mao:
  Mesma regra do resto do projeto -- o arquivo versionado e a fonte da verdade. Mudou um numero no
  BASELINE? Edite aqui e regenere, em vez de cacar a caixa de texto no slide.

Uso:
  python docs/gera_apresentacao.py

Cores e tipografia saem dos tokens do Vitrine DS (Skills/Vitrine Design System/tokens/colors.css).
"""
import pathlib

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

SAIDA = pathlib.Path(__file__).resolve().parent / "apresentacao-vitrine.pptx"

# ---------------------------------------------------------------- tokens do DS
INK_950 = RGBColor(0x0A, 0x12, 0x24)
INK_900 = RGBColor(0x0F, 0x1B, 0x33)
INK_700 = RGBColor(0x2A, 0x3D, 0x63)
INK_500 = RGBColor(0x5B, 0x6B, 0x8C)
INK_400 = RGBColor(0x83, 0x90, 0xAA)
INK_200 = RGBColor(0xCD, 0xD3, 0xE0)
INK_100 = RGBColor(0xE4, 0xE8, 0xF0)
LIME_500 = RGBColor(0xC6, 0xF4, 0x32)
LIME_600 = RGBColor(0xA8, 0xD6, 0x1A)
LIME_100 = RGBColor(0xF2, 0xFC, 0xD0)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PAPER_50 = RGBColor(0xFA, 0xFA, 0xF7)
DANGER = RGBColor(0xE0, 0x47, 0x3C)
WARNING = RGBColor(0xF2, 0xA5, 0x16)
SUCCESS = RGBColor(0x1F, 0x9D, 0x6B)
INFO = RGBColor(0x3B, 0x6F, 0xE0)

FONTE = "Figtree"          # o DS usa so Figtree; cai em sans do sistema se nao instalada
L, T = Inches(0.9), Inches(0.75)
LARGURA_UTIL = Inches(11.5)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


# ------------------------------------------------------------------ utilidades
def slide(fundo):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    r.fill.solid()
    r.fill.fore_color.rgb = fundo
    r.line.fill.background()
    r.shadow.inherit = False
    return s


def caixa(s, x, y, w, h):
    cx = s.shapes.add_textbox(x, y, w, h)
    tf = cx.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return tf


def texto(tf, txt, tamanho, cor, negrito=False, espaco_antes=0, alinhamento=PP_ALIGN.LEFT,
          espacamento=1.0, primeiro=False):
    p = tf.paragraphs[0] if primeiro else tf.add_paragraph()
    p.alignment = alinhamento
    p.space_before = Pt(espaco_antes)
    p.line_spacing = espacamento
    r = p.add_run()
    r.text = txt
    r.font.size = Pt(tamanho)
    r.font.bold = negrito
    r.font.color.rgb = cor
    r.font.name = FONTE
    return p


def wordmark(s, cor_texto=INK_900, x=None, y=None, tamanho=15):
    """`vitrine.` com o ponto em lime -- a marca do DS."""
    x = L if x is None else x
    y = Inches(6.75) if y is None else y
    tf = caixa(s, x, y, Inches(3), Inches(0.4))
    p = tf.paragraphs[0]
    a = p.add_run()
    a.text = "vitrine"
    a.font.size, a.font.bold, a.font.name = Pt(tamanho), True, FONTE
    a.font.color.rgb = cor_texto
    b = p.add_run()
    b.text = "."
    b.font.size, b.font.bold, b.font.name = Pt(tamanho), True, FONTE
    b.font.color.rgb = LIME_500


def regua(s, y, cor=LIME_500, x=None, w=Inches(1.5), h=Pt(5)):
    x = L if x is None else x
    r = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    r.fill.solid()
    r.fill.fore_color.rgb = cor
    r.line.fill.background()
    r.shadow.inherit = False
    return r


def cartao(s, x, y, w, h, fundo=WHITE, borda=None, raio=0.06):
    c = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    c.fill.solid()
    c.fill.fore_color.rgb = fundo
    if borda:
        c.line.color.rgb = borda
        c.line.width = Pt(1.25)
    else:
        c.line.fill.background()
    c.shadow.inherit = False
    try:
        c.adjustments[0] = raio
    except (IndexError, KeyError):
        pass
    return c


def titulo_slide(s, chapeu, titulo, cor_titulo=INK_900, cor_chapeu=INK_500):
    if chapeu:
        tf = caixa(s, L, T, LARGURA_UTIL, Inches(0.3))
        texto(tf, chapeu.upper(), 12, cor_chapeu, True, primeiro=True)
    tf = caixa(s, L, T + Inches(0.34), LARGURA_UTIL, Inches(0.9))
    texto(tf, titulo, 34, cor_titulo, True, espacamento=0.95, primeiro=True)
    regua(s, T + Inches(1.28))


def estatistica(s, x, y, w, numero, rotulo, cor_num=INK_900, fundo=WHITE, borda=INK_100,
                altura=Inches(1.55), tamanho_num=34):
    cartao(s, x, y, w, altura, fundo=fundo, borda=borda)
    tf = caixa(s, x + Inches(0.28), y + Inches(0.22), w - Inches(0.56), altura - Inches(0.4))
    texto(tf, numero, tamanho_num, cor_num, True, espacamento=0.9, primeiro=True)
    texto(tf, rotulo, 11.5, INK_500, espaco_antes=5)


def tabela(s, x, y, w, colunas, linhas, larguras, altura_linha=Inches(0.42)):
    n_l, n_c = len(linhas) + 1, len(colunas)
    forma = s.shapes.add_table(n_l, n_c, x, y, w, altura_linha * n_l)
    t = forma.table
    t.first_row = False
    for i, lg in enumerate(larguras):
        t.columns[i].width = Emu(int(w * lg))
    for j, c in enumerate(colunas):
        cel = t.cell(0, j)
        cel.text = c
        cel.fill.solid()
        cel.fill.fore_color.rgb = INK_900
        cel.vertical_anchor = MSO_ANCHOR.MIDDLE
        cel.margin_left = cel.margin_right = Inches(0.12)
        p = cel.text_frame.paragraphs[0]
        p.runs[0].font.size, p.runs[0].font.bold = Pt(11), True
        p.runs[0].font.color.rgb = WHITE
        p.runs[0].font.name = FONTE
    for i, linha in enumerate(linhas, start=1):
        for j, v in enumerate(linha):
            cel = t.cell(i, j)
            cel.text = str(v)
            cel.fill.solid()
            cel.fill.fore_color.rgb = WHITE if i % 2 else PAPER_50
            cel.vertical_anchor = MSO_ANCHOR.MIDDLE
            cel.margin_left = cel.margin_right = Inches(0.12)
            p = cel.text_frame.paragraphs[0]
            destaque = str(v).startswith("**")
            if destaque:
                p.runs[0].text = str(v).replace("**", "")
            p.runs[0].font.size = Pt(10.5)
            p.runs[0].font.bold = destaque
            p.runs[0].font.color.rgb = INK_900 if destaque else INK_700
            p.runs[0].font.name = FONTE
    return t


def rodape(s, txt, cor=INK_400):
    tf = caixa(s, L, Inches(6.78), LARGURA_UTIL, Inches(0.3))
    texto(tf, txt, 9.5, cor, primeiro=True)


# ============================================================ 1 · capa
s = slide(INK_900)
regua(s, Inches(2.15), LIME_500, w=Inches(2.2), h=Pt(7))
tf = caixa(s, L, Inches(2.45), Inches(10.5), Inches(2.2))
p = tf.paragraphs[0]
a = p.add_run(); a.text = "vitrine"
a.font.size, a.font.bold, a.font.name, a.font.color.rgb = Pt(72), True, FONTE, WHITE
b = p.add_run(); b.text = "."
b.font.size, b.font.bold, b.font.name, b.font.color.rgb = Pt(72), True, FONTE, LIME_500
texto(tf, "Do dado cru à pergunta em português", 34, WHITE, espaco_antes=14, espacamento=1.0)
texto(tf, "Lakehouse bronze → silver → gold, três dashboards AI/BI e um agente do Genie "
          "que responde à diretoria sem ninguém escrever SQL.", 15, INK_200, espaco_antes=14)
tf = caixa(s, L, Inches(6.35), Inches(11), Inches(0.5))
texto(tf, "Databricks  ·  Lakeflow Declarative Pipelines  ·  Unity Catalog  ·  AI/BI  ·  Genie"
          "        período dos dados: 13/12/2025 a 11/01/2026", 11, INK_400, primeiro=True)

# ============================================================ 2 · o desafio
s = slide(PAPER_50)
titulo_slide(s, "o ponto de partida", "Três diretorias, uma base, zero autonomia")
tf = caixa(s, L, Inches(2.25), Inches(5.4), Inches(3.2))
texto(tf, "Comercial, Customer Success e Pricing precisavam das mesmas tabelas para perguntas "
          "completamente diferentes — e dependiam de alguém escrever SQL para cada pergunta.",
      15, INK_700, espacamento=1.25, primeiro=True)
texto(tf, "A base ainda trazia problemas conhecidos: venda de produto fora do catálogo, venda "
          "anterior ao cadastro do produto e cotação de concorrente pela metade do nosso preço.",
      15, INK_700, espaco_antes=14, espacamento=1.25)
texto(tf, "Apagar linha suja resolveria o relatório e destruiria a receita.", 15, INK_900, True,
      espaco_antes=14, espacamento=1.25)
for i, (num, rot, cor) in enumerate([
    ("3.020", "vendas no período", INK_900),
    ("R$ 974.077,28", "receita a fechar em toda camada", INK_900),
    ("80", "linhas com problema conhecido, todas preservadas", DANGER),
    ("0", "pergunta respondida sem um analista", DANGER),
]):
    x = Inches(6.6) + Inches(2.72) * (i % 2)
    y = Inches(2.25) + Inches(1.75) * (i // 2)
    estatistica(s, x, y, Inches(2.5), num, rot, cor_num=cor)
wordmark(s)

# ============================================================ 3 · andamento
s = slide(PAPER_50)
titulo_slide(s, "andamento", "Como o projeto foi construído")
etapas = [
    ("01", "Silver", "Limpa, deduplica e MARCA o problema.\nNunca descarta linha."),
    ("02", "Gold", "5 tabelas, uma por diretoria.\n70 colunas com COMMENT."),
    ("03", "Dashboards", "3 painéis AI/BI como código,\nno Vitrine Design System."),
    ("04", "Genie", "Um agente para as três\ndiretorias, em português."),
    ("05", "Placar", "10/10 + 2/2, em 4 rodadas\nde ajuste de contexto."),
    ("06", "Blindagem", "4 redes que impedem\na regressão silenciosa."),
]
for i, (n, t, d) in enumerate(etapas):
    x = L + Inches(3.85) * (i % 3)
    y = Inches(2.3) + Inches(2.0) * (i // 3)
    cartao(s, x, y, Inches(3.55), Inches(1.72), borda=INK_100)
    tf = caixa(s, x + Inches(0.3), y + Inches(0.24), Inches(2.9), Inches(1.3))
    p = texto(tf, n, 13, LIME_600, True, primeiro=True)
    texto(tf, t, 19, INK_900, True, espaco_antes=2)
    texto(tf, d, 11.5, INK_500, espaco_antes=6, espacamento=1.15)
wordmark(s)

# ============================================================ 4 · arquitetura
s = slide(INK_900)
titulo_slide(s, "arquitetura", "Um pipeline. Quatro destinos.", cor_titulo=WHITE,
             cor_chapeu=INK_400)
camadas = [
    ("bronze", "vendas · produtos\nclientes · concorrentes", INK_700, INK_200,
     "sobrescrita por outro processo"),
    ("silver", "limpa, deduplica\ne MARCA problemas", INK_500, WHITE, "4 tabelas · @dp.expect"),
    ("gold", "uma tabela\npor diretoria", LIME_500, INK_900, "5 tabelas · 70 colunas"),
    ("consumo", "3 dashboards\n+ Genie", WHITE, INK_900, "Ask Genie nos três"),
]
for i, (nome, corpo, fundo, cor_txt, nota) in enumerate(camadas):
    x = L + Inches(2.95) * i
    cartao(s, x, Inches(2.5), Inches(2.6), Inches(2.1), fundo=fundo)
    tf = caixa(s, x + Inches(0.28), Inches(2.72), Inches(2.05), Inches(1.7))
    texto(tf, nome, 20, cor_txt, True, primeiro=True)
    texto(tf, corpo, 12, cor_txt, espaco_antes=8, espacamento=1.2)
    tf = caixa(s, x, Inches(4.75), Inches(2.6), Inches(0.5))
    texto(tf, nota, 10, INK_400, primeiro=True)
    if i < 3:
        tf = caixa(s, x + Inches(2.63), Inches(3.35), Inches(0.35), Inches(0.4))
        texto(tf, "→", 20, LIME_500, True, alinhamento=PP_ALIGN.CENTER, primeiro=True)
tf = caixa(s, L, Inches(5.5), Inches(11.5), Inches(1.0))
texto(tf, "Tudo em um único Lakeflow Declarative Pipeline serverless. As dependências são "
          "implícitas: o pipeline monta o grafo pelo nome da tabela no código.", 14, INK_200,
      espacamento=1.25, primeiro=True)
texto(tf, "Materialized view em batch, nunca streaming — a bronze é sobrescrita fora do nosso "
          "controle, e streaming exige fonte append-only.", 12, INK_400, espaco_antes=8)
wordmark(s, cor_texto=WHITE)

# ============================================================ 5 · silver
s = slide(PAPER_50)
titulo_slide(s, "camada silver", "Problema de qualidade se marca. Nunca se descarta.")
tf = caixa(s, L, Inches(2.2), Inches(11.5), Inches(0.5))
texto(tf, "Apagar uma venda suja é apagar receita. A silver marca o problema numa coluna booleana "
          "e o mede com @dp.expect — a linha continua contando no faturamento.", 14, INK_700,
      espacamento=1.2, primeiro=True)
tabela(s, L, Inches(3.0), Inches(11.5),
       ["problema conhecido", "volume", "como fica marcado"],
       [["Venda de produto fora do catálogo", "**20 vendas · R$ 4.240,01",
         "produto_cadastrado = false"],
        ["Venda anterior ao cadastro do produto", "**5 vendas · R$ 325,88",
         "venda_antes_do_cadastro = true"],
        ["Cotação de concorrente suspeita", "**55 cotações · 15 produtos",
         "possui_preco_suspeito = true"],
        ["Nome de cliente com pronome de tratamento", "**11 clientes",
         "corrigido, original preservado"]],
       [0.42, 0.29, 0.29])
tf = caixa(s, L, Inches(5.5), Inches(11.5), Inches(0.9))
texto(tf, "@dp.expect conta as linhas saudáveis: a expectation afirma o estado bom, e a métrica de "
          "falha é exatamente a contagem do problema.", 13, INK_900, True, primeiro=True)
texto(tf, "Expectation vermelha esperada não é regressão — o BASELINE.md diz qual é o número certo "
          "de cada uma.", 12, INK_500, espaco_antes=6)
wordmark(s)

# ============================================================ 6 · gold
s = slide(PAPER_50)
titulo_slide(s, "camada gold", "Escrita para uma IA ler")
tf = caixa(s, L, Inches(2.2), Inches(6.0), Inches(3.4))
texto(tf, "As 70 colunas carregam tipo e COMMENT com unidade, regra de cálculo e a armadilha que "
          "evita erro: o que significa zero, o que significa nulo, e o que a coluna NÃO é.",
      14.5, INK_700, espacamento=1.25, primeiro=True)
texto(tf, "Sem o tipo declarado, o Databricks descarta o comentário em silêncio.", 13, INK_900,
      True, espaco_antes=12)
texto(tf, "É esse texto que o Genie lê para escrever SQL.\nComentário vago = resposta adivinhada.",
      14.5, INK_900, True, espaco_antes=12, espacamento=1.2)
cartao(s, Inches(7.3), Inches(2.2), Inches(5.1), Inches(2.5), fundo=INK_900)
tf = caixa(s, Inches(7.6), Inches(2.45), Inches(4.5), Inches(2.1))
texto(tf, "ticket_medio  DECIMAL(10,2)", 12.5, LIME_500, True, primeiro=True)
texto(tf, "\"Valor médio de UMA venda, em reais (R$).\nCálculo: receita total dividida pelo número "
          "de vendas — nunca média de médias.\nNÃO é o preço unitário.\"", 12, WHITE,
      espaco_antes=8, espacamento=1.3)
for i, (n, r) in enumerate([("5", "tabelas gold"), ("70", "colunas comentadas"),
                            ("0", "colunas sem comentário")]):
    estatistica(s, L + Inches(4.0) * i, Inches(5.05), Inches(3.7), n, r, altura=Inches(1.3),
                tamanho_num=30)
wordmark(s)

# ============================================================ 7 · dashboards
s = slide(PAPER_50)
titulo_slide(s, "produto", "Três dashboards, três perguntas de diretoria")
painel = [
    ("Comercial", "Quanto vendemos, quando,\nem qual canal e com quais produtos?",
     "Quarta vende mais que sábado:\no período tem 5 sábados e só 4 quartas."),
    ("Customer Success", "Quem são os melhores clientes\ne onde eles estão?",
     "O maior bloco não é o VIP:\nTOP_TIER faz 50,6% da receita."),
    ("Pricing", "Estamos mais caros que o mercado,\ne em quais produtos agir?",
     "Sem problema sistêmico: o risco real\nestá em 20 produtos — 71,8% em um só."),
]
for i, (nome, pergunta, achado) in enumerate(painel):
    x = L + Inches(3.9) * i
    cartao(s, x, Inches(2.25), Inches(3.6), Inches(3.5), borda=INK_100)
    tf = caixa(s, x + Inches(0.3), Inches(2.5), Inches(3.0), Inches(3.0))
    texto(tf, nome, 20, INK_900, True, primeiro=True)
    texto(tf, pergunta, 12, INK_500, espaco_antes=8, espacamento=1.25)
    p = texto(tf, "o que ele já entregou", 10, LIME_600, True, espaco_antes=16)
    texto(tf, achado, 12.5, INK_900, True, espaco_antes=6, espacamento=1.25)
tf = caixa(s, L, Inches(6.0), Inches(11.5), Inches(0.6))
texto(tf, "Os três seguem o Vitrine Design System — e a identidade é cobrada por validador, "
          "não confiada à memória de quem edita.", 13, INK_700, primeiro=True)
wordmark(s)

# ============================================================ 8 · design system
s = slide(INK_900)
titulo_slide(s, "identidade", "A marca não é enfeite. É contrato.", cor_titulo=WHITE,
             cor_chapeu=INK_400)
swatches = [("ink 900", INK_900, WHITE), ("ink 700", INK_700, WHITE), ("ink 400", INK_400, INK_900),
            ("lime 500", LIME_500, INK_900), ("âmbar", WARNING, INK_900), ("info", INFO, WHITE)]
for i, (nome, cor, txt) in enumerate(swatches):
    x = L + Inches(1.85) * i
    c = cartao(s, x, Inches(2.35), Inches(1.65), Inches(1.15), fundo=cor)
    if cor == INK_900:
        c.line.color.rgb = INK_700
        c.line.width = Pt(1)
    tf = caixa(s, x + Inches(0.18), Inches(3.05), Inches(1.4), Inches(0.35))
    texto(tf, nome, 10, txt, True, primeiro=True)
itens = [
    ("Figtree em tudo", "a única fonte do DS, aplicada pelo tema do dashboard"),
    ("Cantos de card em 20px", "o DS usa sombra, não traço: borda igual ao fundo"),
    ("Lime é acento, não enfeite", "seleção, filtro ativo e o segmento que o CS persegue"),
    ("Copy em português de verdade", "com acento, sentence case, sem CAIXA ALTA, tratando por você"),
]
for i, (t, d) in enumerate(itens):
    x = L + Inches(5.75) * (i % 2)
    y = Inches(3.9) + Inches(0.95) * (i // 2)
    tf = caixa(s, x, y, Inches(5.3), Inches(0.8))
    texto(tf, t, 15, LIME_500, True, primeiro=True)
    texto(tf, d, 12, INK_200, espaco_antes=3, espacamento=1.15)
tf = caixa(s, L, Inches(6.0), Inches(11.5), Inches(0.5))
texto(tf, "O validador recusa cor fora dos tokens, tema divergente entre painéis e texto sem "
          "acento. Regressão de marca não passa despercebida.", 12.5, WHITE, True, primeiro=True)
wordmark(s, cor_texto=WHITE)

# ============================================================ 9 · genie (cereja)
s = slide(INK_950)
regua(s, Inches(1.5), LIME_500, w=Inches(2.2), h=Pt(7))
tf = caixa(s, L, Inches(1.85), Inches(7.6), Inches(3.6))
texto(tf, "A CEREJA DO BOLO", 13, LIME_500, True, primeiro=True)
texto(tf, "O Genie", 58, WHITE, True, espaco_antes=8, espacamento=0.95)
texto(tf, "Um único agente atende as três diretorias. Ele sabe apenas o que o projeto ensinou: "
          "as 5 tabelas gold, 2.481 caracteres de regra de negócio, 4 SQL de exemplo para as "
          "contas que a IA erra e 6 perguntas de partida.", 15, INK_200, espaco_antes=14,
      espacamento=1.3)
texto(tf, "A diretoria pergunta em português. Em segundos.", 18, LIME_500, True, espaco_antes=14)
for i, (t, d) in enumerate([
    ("Só a gold", "nenhuma tabela bronze ou silver entra no space"),
    ("Joins e sinônimos escritos à mão", "\"faturamento\" é receita, \"UF\" é estado"),
    ("SQL de exemplo onde a IA erra", "ticket médio, participação, contagem de preço"),
    ("6 perguntas de partida", "todas testadas: nenhuma volta vazia"),
]):
    y = Inches(4.72) + Inches(0.48) * i
    regua(s, y + Inches(0.07), LIME_500, x=L, w=Pt(5), h=Pt(13))
    tf2 = caixa(s, L + Inches(0.3), y, Inches(7.4), Inches(0.45))
    p = tf2.paragraphs[0]
    a = p.add_run(); a.text = t + "  "
    a.font.size, a.font.bold, a.font.name, a.font.color.rgb = Pt(13), True, FONTE, WHITE
    b = p.add_run(); b.text = d
    b.font.size, b.font.name, b.font.color.rgb = Pt(12), FONTE, INK_400
for i, (n, r) in enumerate([("5", "tabelas gold,\nnenhuma silver"),
                            ("2.481", "caracteres de\ninstrução"),
                            ("12", "perguntas no\nteste de aceitação")]):
    y = Inches(1.9) + Inches(1.6) * i
    cartao(s, Inches(9.1), y, Inches(3.3), Inches(1.35), fundo=INK_900)
    tf = caixa(s, Inches(9.4), y + Inches(0.2), Inches(2.8), Inches(1.0))
    texto(tf, n, 30, LIME_500, True, espacamento=0.9, primeiro=True)
    texto(tf, r, 11, INK_200, espaco_antes=3, espacamento=1.15)
wordmark(s, cor_texto=WHITE)

# ============================================================ 10 · placar
s = slide(PAPER_50)
titulo_slide(s, "prova", "O placar: 10/10, em quatro rodadas")
tf = caixa(s, L, Inches(2.2), Inches(11.5), Inches(0.5))
texto(tf, "Uma conversa nova por pergunta (conversa reaproveitada esconde regressão), resposta "
          "comparada com SQL rodado na gold na hora. Quando errou, mudou o contexto — nunca o "
          "esperado.", 13.5, INK_700, espacamento=1.2, primeiro=True)
tabela(s, L, Inches(3.05), Inches(11.5), ["rodada", "placar", "o que mudou entre elas"],
       [["1", "2/10", "Era bug do TESTE: o gabarito devolve 974077.28 e o Genie responde "
                      "R$ 974.077,28"],
        ["2", "8/10", "Gabarito sem a coluna de total; e a regra de clientes era permissão, "
                      "não obrigação"],
        ["3", "9/10", "Regra virou imperativa no texto — e ainda assim falhou"],
        ["**4", "**10/10 + 2/2", "**A regra entrou no SQL de exemplo. Regra que falha em texto "
                                 "passa como exemplo."]],
       [0.10, 0.16, 0.74])
tf = caixa(s, L, Inches(5.45), Inches(11.5), Inches(1.0))
texto(tf, "A lição que mais ensinou: desconfie do teste antes do modelo.", 15, INK_900, True,
      primeiro=True)
texto(tf, "O 2/10 da primeira rodada era formatação de número no script de placar — o agente já "
          "acertava 8 das 10. Rodada 5 repetiu 10/10 sem nenhuma mudança, descartando sorte.",
      12.5, INK_500, espaco_antes=6, espacamento=1.2)
wordmark(s)

# ============================================================ 11 · dizer não
s = slide(PAPER_50)
titulo_slide(s, "confiança", "O que o torna confiável é ele saber dizer não")
for i, (pergunta, resposta, cor) in enumerate([
    ("\"Qual foi o nosso lucro?\"",
     "\"Não há dados de custo, margem, lucro, desconto ou imposto neste banco; só existe receita "
     "bruta. Não é possível calcular o lucro com as informações disponíveis.\"", DANGER),
    ("\"Quanto vendemos ontem?\"",
     "\"Não é possível responder perguntas sobre 'ontem', pois os dados cobrem apenas o período de "
     "13/12/2025 a 11/01/2026. O último dia com vendas é 11/01/2026; deseja ver esse dia?\"",
     WARNING),
]):
    y = Inches(2.3) + Inches(1.85) * i
    cartao(s, L, y, Inches(11.5), Inches(1.6), borda=INK_100)
    regua(s, y + Inches(0.25), cor, x=L + Inches(0.3), w=Pt(5), h=Inches(1.1))
    tf = caixa(s, L + Inches(0.65), y + Inches(0.25), Inches(10.5), Inches(1.2))
    texto(tf, pergunta, 15, INK_900, True, primeiro=True)
    texto(tf, resposta, 12.5, INK_700, espaco_antes=7, espacamento=1.25)
tf = caixa(s, L, Inches(6.05), Inches(11.5), Inches(0.6))
texto(tf, "Na pergunta sobre \"ontem\" ele não gera SQL nenhum: explica o período e oferece o "
          "último dia disponível. Um agente que inventa número perde a diretoria na primeira "
          "resposta.", 13, INK_900, True, espacamento=1.2, primeiro=True)
wordmark(s)

# ============================================================ 12 · impacto
s = slide(INK_900)
titulo_slide(s, "impacto", "O que essas respostas colocam sobre a mesa",
             cor_titulo=WHITE, cor_chapeu=INK_400)
cartao(s, L, Inches(2.25), Inches(5.5), Inches(2.3), fundo=LIME_500)
tf = caixa(s, L + Inches(0.35), Inches(2.5), Inches(4.8), Inches(1.9))
texto(tf, "R$ 654.436,22", 42, INK_900, True, espacamento=0.9, primeiro=True)
texto(tf, "de receita colocada sob decisão explícita — 67,2% do faturamento do período, que antes "
          "dependia de alguém pensar na pergunta certa.", 12.5, INK_900, espaco_antes=8,
      espacamento=1.25)
blocos = [
    ("R$ 161.375,09", "exposta à concorrência: produtos confirmadamente mais caros que TODOS os "
                      "rivais", WARNING),
    ("71,8% em 1 produto", "a Camisa Social sozinha responde por R$ 115.794,92 — um ajuste resolve "
                           "três quartos do risco", LIME_500),
    ("R$ 493.061,13", "no bloco TOP_TIER (50,6%), invisível para uma estratégia focada só nos VIP",
     INFO),
    ("15 produtos", "\"mais caros que todos\" que eram bug de coleta: corte de preço errado evitado",
     DANGER),
]
for i, (n, d, cor) in enumerate(blocos):
    x = Inches(6.8)
    y = Inches(2.25) + Inches(1.12) * i
    tf = caixa(s, x, y, Inches(5.6), Inches(1.0))
    texto(tf, n, 19, cor, True, primeiro=True)
    texto(tf, d, 11.5, INK_200, espaco_antes=2, espacamento=1.15)
tf = caixa(s, L, Inches(4.8), Inches(5.5), Inches(1.5))
texto(tf, "Nenhum desses números foi estimado: todos saem de consulta na gold, conferida contra o "
          "BASELINE. O que muda com o Genie é quem consegue chegar até eles — e em quanto tempo.",
      12.5, INK_200, espacamento=1.25, primeiro=True)
rodape(s, "Marca e dados são de estudo (Vitrine é fictícia). Período: 13/12/2025 a 11/01/2026.")
wordmark(s, cor_texto=WHITE, y=Inches(6.4))

# ============================================================ 13 · blindagem
s = slide(PAPER_50)
titulo_slide(s, "engenharia", "Quatro redes contra a regressão silenciosa")
tabela(s, L, Inches(2.35), Inches(11.5), ["rede", "o que pega", "quando"],
       [["@dp.expect", "problema dentro de UMA tabela, durante a escrita", "no pipeline"],
        ["testes_qualidade.py", "**18 testes cruzando tabelas, sobre o que já foi publicado",
         "task 2 do job"],
        ["valida_dashboards.py", "o que o bundle validate NÃO valida: measure fantasma, grid "
                                 "sobreposto, tema fora do DS, copy sem acento",
         "antes do deploy"],
        ["valida_genie.py", "schema do space, ordem exigida pelo servidor, teto de caracteres, "
                            "pergunta que entrega a resposta do teste", "antes do deploy"]],
       [0.22, 0.58, 0.20], altura_linha=Inches(0.62))
cartao(s, L, Inches(5.2), Inches(11.5), Inches(1.25), fundo=INK_900)
tf = caixa(s, L + Inches(0.4), Inches(5.42), Inches(10.7), Inches(0.9))
texto(tf, "\"Deploy passou\" não é evidência de nada.", 17, LIME_500, True, primeiro=True)
texto(tf, "Um .lvdash.json com dataset errado deploya limpo e quebra só no navegador. Cada regra "
          "dentro dos validadores é a cicatriz de um erro que já aconteceu aqui.", 12.5, WHITE,
      espaco_antes=6, espacamento=1.2)
wordmark(s)

# ============================================================ 14 · fecho
s = slide(INK_900)
regua(s, Inches(2.4), LIME_500, w=Inches(2.2), h=Pt(7))
tf = caixa(s, L, Inches(2.75), Inches(10.6), Inches(3.0))
texto(tf, "Dado confiável, com marca e com voz.", 44, WHITE, True, espacamento=1.0, primeiro=True)
texto(tf, "A diretoria não abre uma tabela: ela faz uma pergunta. E recebe o número certo — ou um "
          "\"esse dado não existe aqui\", que vale igual.", 17, INK_200, espaco_antes=16,
      espacamento=1.3)
tf = caixa(s, L, Inches(5.5), Inches(11), Inches(0.5))
texto(tf, "10/10 + 2/2 no placar  ·  18 testes de dados  ·  70 colunas documentadas  ·  "
          "4 redes de proteção", 13, LIME_500, True, primeiro=True)
wordmark(s, cor_texto=WHITE, tamanho=20)

prs.save(SAIDA)
print(f"{SAIDA.name}: {len(prs.slides.__iter__.__self__._sldIdLst)} slides")
